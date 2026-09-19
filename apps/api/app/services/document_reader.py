import io
import os
import time
import zipfile
from pathlib import Path

import requests
from docx import Document
from sarvamai import SarvamAI
from sarvamai.core.api_error import ApiError

SUPPORTED = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".txt",
    ".docx",
    ".zip",
}

client = SarvamAI(
    api_subscription_key=os.getenv("SARVAM_API_KEY")
)

TERMINAL_STATES = {
    "completed",
    "partially_completed",
    "failed",
    "rejected",
}


def normalize_text(value: str) -> str:
    lines = []
    seen = set()

    for raw in value.replace("\x00", " ").splitlines():
        line = " ".join(raw.split()).strip()
        if not line:
            continue

        key = line.lower()

        if key not in seen:
            seen.add(key)
            lines.append(line)

    return "\n".join(lines).strip()


def read_txt(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return normalize_text(data.decode(encoding))
        except UnicodeDecodeError:
            continue

    return ""


def read_docx_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))

    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                parts.append(" | ".join(values))

    return normalize_text("\n".join(parts))


def extract_docx_images(data: bytes) -> list[tuple[str, bytes]]:
    images = []

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for name in archive.namelist():
            if not name.startswith("word/media/"):
                continue

            suffix = Path(name).suffix.lower()

            if suffix not in {".png", ".jpg", ".jpeg"}:
                continue

            images.append((Path(name).name, archive.read(name)))

            if len(images) >= 10:
                break

    return images


def build_image_zip(images: list[tuple[str, bytes]]) -> bytes:
    output = io.BytesIO()

    with zipfile.ZipFile(
        output,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:
        for filename, data in images[:10]:
            archive.writestr(filename, data)

    return output.getvalue()


def download_digitised_output(job_id: str) -> str:
    download = client.doc_ai.get_download_url(job_id=job_id)

    response = requests.get(
        download.url,
        timeout=120,
    )

    response.raise_for_status()

    archive = zipfile.ZipFile(io.BytesIO(response.content))

    markdown_files = [
        name
        for name in archive.namelist()
        if name.lower().endswith((".md", ".markdown", ".txt"))
    ]

    parts = []

    for name in sorted(markdown_files):
        parts.append(
            archive.read(name).decode(
                "utf-8",
                errors="replace",
            )
        )

    return normalize_text("\n".join(parts))


def sarvam_ocr(
    filename: str,
    data: bytes,
    content_type: str | None = None,
) -> str:
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":
        mime = "application/pdf"
    elif suffix == ".png":
        mime = "image/png"
    elif suffix in {".jpg", ".jpeg"}:
        mime = "image/jpeg"
    elif suffix == ".zip":
        mime = "application/zip"
    else:
        raise ValueError(
            f"Sarvam OCR does not support {suffix or 'this file type'}."
        )

    stream = io.BytesIO(data)

    try:
        job = client.doc_ai.digitise(
            file=[
                (
                    filename,
                    stream,
                    content_type or mime,
                )
            ],
            language="en-IN",
            output_format="md",
        )

        for _ in range(120):
            status = client.doc_ai.get_status(
                job_id=job.job_id
            )

            state = status.status.lower()

            if state in TERMINAL_STATES:
                break

            time.sleep(2)

        if state not in {"completed", "partially_completed"}:
            raise ValueError(
                f"Sarvam Document AI failed with status: {state}"
            )

        return download_digitised_output(job.job_id)

    except ApiError as exc:
        raise ValueError(
            f"Sarvam Document AI error {exc.status_code}: {exc.body}"
        ) from exc


def extract_resume_text(
    filename: str,
    data: bytes,
    content_type: str | None = None,
) -> str:
    suffix = Path(filename).suffix.lower()

    if suffix not in SUPPORTED:
        raise ValueError(
            "Unsupported resume format. Use PDF, DOCX, TXT, PNG, JPG, JPEG, or ZIP."
        )

    if not data:
        raise ValueError("The uploaded resume is empty.")

    if suffix == ".txt":
        text = read_txt(data)

        if len(text) >= 120:
            return text

        raise ValueError(
            "The TXT resume does not contain enough readable content."
        )

    if suffix == ".docx":
        text = read_docx_text(data)
        images = extract_docx_images(data)

        ocr_text = ""

        if images:
            image_zip = build_image_zip(images)

            ocr_text = sarvam_ocr(
                "resume-images.zip",
                image_zip,
                "application/zip",
            )

        combined = normalize_text(
            f"{text}\n{ocr_text}"
        )

        if len(combined) < 120:
            raise ValueError(
                "JOBIX could not extract enough readable resume content from the DOCX."
            )

        return combined

    if suffix == ".zip":
        text = sarvam_ocr(
            filename,
            data,
            "application/zip",
        )

        if len(text) < 120:
            raise ValueError(
                "JOBIX could not extract enough readable content from the ZIP."
            )

        return text

    text = sarvam_ocr(
        filename,
        data,
        content_type,
    )

    if len(text) < 120:
        raise ValueError(
            "JOBIX could not extract enough readable resume content from the uploaded document."
        )

    return text
