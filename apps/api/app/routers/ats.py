import asyncio
from typing import Any
import json
import re
from json_repair import repair_json
from io import BytesIO

import requests
from docx import Document
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from pydantic import BaseModel, Field
from pypdf import PdfReader
import fitz

from app.services.document_reader import extract_resume_text as extract_uploaded_resume
from app.core.config import settings
from app.core.security import get_current_user
from app.models.user import User

router = APIRouter(prefix="/ats", tags=["ATS"])

MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_RESUME_TEXT = 14000
MAX_JOB_DESCRIPTION = 6000


class AtsItem(BaseModel):
    title: str
    explanation: str


class AtsImprovement(BaseModel):
    title: str
    explanation: str
    priority: str


class AtsResult(BaseModel):
    ats_score: int = Field(ge=0, le=100)
    score_label: str
    role_match_score: int = Field(ge=0, le=100)
    role: str
    summary: str
    strengths: list[AtsItem]
    missing_fields: list[AtsItem]
    improvements: list[AtsImprovement]
    keyword_gaps: list[str]
    matched_keywords: list[str]
    formatting_issues: list[str]
    action_plan: list[str]


def extract_pdf(file: UploadFile) -> str:
    try:
        file.file.seek(0)

        reader = PdfReader(file.file)
        parts = []

        for page in reader.pages:
            text = page.extract_text() or ""

            if text.strip():
                parts.append(text)

        text = "\n".join(parts).strip()

        if len(text) >= 120:
            return text

        file.file.seek(0)

        pdf_bytes = file.file.read()
        document = fitz.open(
            stream=pdf_bytes,
            filetype="pdf",
        )

        fallback_parts = []

        for page in document:
            page_text = page.get_text(
                "text",
                sort=True,
            )

            if page_text.strip():
                fallback_parts.append(page_text)

        document.close()

        fallback_text = "\n".join(
            fallback_parts
        ).strip()

        if len(fallback_text) > len(text):
            return fallback_text

        return text

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read the PDF: {exc}",
        )

def extract_docx(file: UploadFile) -> str:
    try:
        document = Document(file.file)
        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        return "\n".join(parts)
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read the DOCX: {exc}",
        )


def extract_txt(file: UploadFile) -> str:
    try:
        return file.file.read().decode(
            "utf-8",
            errors="ignore",
        )
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read the text file: {exc}",
        )



def build_schema() -> dict[str, Any]:
    item_schema = {
        "type": "object",
        "properties": {
            "title": {"type": "string", "maxLength": 80},
            "explanation": {"type": "string", "maxLength": 180},
        },
        "required": ["title", "explanation"],
        "additionalProperties": False,
    }

    improvement_schema = {
        "type": "object",
        "properties": {
            "title": {"type": "string", "maxLength": 80},
            "explanation": {"type": "string", "maxLength": 180},
            "priority": {"type": "string", "maxLength": 20},
        },
        "required": ["title", "explanation", "priority"],
        "additionalProperties": False,
    }

    return {
        "type": "object",
        "properties": {
            "ats_score": {
                "type": "integer",
                "minimum": 0,
                "maximum": 100,
            },
            "score_label": {
                "type": "string",
                "maxLength": 30,
            },
            "role_match_score": {
                "type": "integer",
                "minimum": 0,
                "maximum": 100,
            },
            "role": {
                "type": "string",
                "maxLength": 80,
            },
            "summary": {
                "type": "string",
                "maxLength": 300,
            },
            "strengths": {
                "type": "array",
                "maxItems": 5,
                "items": item_schema,
            },
            "missing_fields": {
                "type": "array",
                "maxItems": 4,
                "items": item_schema,
            },
            "improvements": {
                "type": "array",
                "maxItems": 6,
                "items": improvement_schema,
            },
            "keyword_gaps": {
                "type": "array",
                "maxItems": 6,
                "items": {
                    "type": "string",
                    "maxLength": 50,
                },
            },
            "matched_keywords": {
                "type": "array",
                "maxItems": 18,
                "items": {
                    "type": "string",
                    "maxLength": 50,
                },
            },
            "formatting_issues": {
                "type": "array",
                "maxItems": 5,
                "items": {
                    "type": "string",
                    "maxLength": 120,
                },
            },
            "action_plan": {
                "type": "array",
                "maxItems": 6,
                "items": {
                    "type": "string",
                    "maxLength": 120,
                },
            },
        },
        "required": [
            "ats_score",
            "score_label",
            "role_match_score",
            "role",
            "summary",
            "strengths",
            "missing_fields",
            "improvements",
            "keyword_gaps",
            "matched_keywords",
            "formatting_issues",
            "action_plan",
        ],
        "additionalProperties": False,
    }

def build_prompt(
    resume_text: str,
    job_role: str,
    job_description: str,
) -> str:
    role = job_role.strip() or "Infer the best target role from the resume."
    jd = (
        job_description.strip()
        or "No job description supplied. Evaluate against the target role."
    )

    return f"""
You are JOBIX ATS.

Analyze the resume for the target role.

TARGET ROLE:
{role}

JOB DESCRIPTION:
{jd}

RESUME:
{resume_text}

Evaluate:
- ATS readability
- role match
- job description match
- skills
- experience
- projects
- education
- certifications
- achievements
- keywords
- missing information
- formatting
- recruiter readability

Rules:
- Use only information actually present in the resume.
- Never invent employers, experience, skills, certifications, metrics or technologies.
- matched_keywords must exist in the resume.
- keyword_gaps must relate to the target role or job description.
- missing_fields must identify genuinely missing or weak information.
- strengths must be supported by the resume.
- improvements must be practical.
- Scores must be realistic.
- Keep summary concise.
- Keep explanations concise.
- Return every required field.
- Empty arrays are allowed.
- Return ONLY the requested JSON object.
""".strip()

def parse_content(content: Any) -> str:
    if content is None:
        return ""

    if isinstance(content, str):
        return content.strip()

    if isinstance(content, dict):
        if isinstance(content.get("text"), str):
            return content["text"].strip()
        return json.dumps(content)

    if isinstance(content, list):
        parts = []

        for item in content:
            if isinstance(item, str):
                parts.append(item)
            elif isinstance(item, dict):
                text = item.get("text")
                if isinstance(text, str):
                    parts.append(text)

        return "".join(parts).strip()

    return str(content).strip()


def parse_json_text(text: str) -> dict[str, Any]:
    cleaned = text.strip()

    if not cleaned:
        raise ValueError("Sarvam returned an empty response.")

    if cleaned.startswith("```"):
        lines = cleaned.splitlines()

        if lines and lines[0].strip().startswith("```"):
            lines = lines[1:]

        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]

        cleaned = "\n".join(lines).strip()

    try:
        parsed = json.loads(cleaned)

        if isinstance(parsed, dict):
            return parsed
    except json.JSONDecodeError:
        pass

    first = cleaned.find("{")
    last = cleaned.rfind("}")

    if first >= 0 and last > first:
        candidate = cleaned[first:last + 1]

        try:
            parsed = json.loads(candidate)

            if isinstance(parsed, dict):
                return parsed
        except json.JSONDecodeError:
            pass

        try:
            repaired = repair_json(candidate)
            parsed = json.loads(repaired)

            if isinstance(parsed, dict):
                return parsed
        except Exception:
            pass

    try:
        repaired = repair_json(cleaned)
        parsed = json.loads(repaired)

        if isinstance(parsed, dict):
            return parsed
    except Exception:
        pass

    raise ValueError("Sarvam returned malformed JSON.")

def normalize_result(
    data: dict[str, Any],
    job_role: str,
) -> AtsResult:

    def string(
        value: Any,
        fallback: str = "",
    ) -> str:
        if value is None:
            return fallback
        return str(value).strip()

    def integer(
        value: Any,
        fallback: int = 0,
    ) -> int:
        try:
            return max(
                0,
                min(
                    100,
                    int(float(value)),
                ),
            )
        except (TypeError, ValueError):
            return fallback

    def string_list(
        value: Any,
    ) -> list[str]:
        if not isinstance(value, list):
            return []

        result = []

        for item in value:
            if isinstance(item, str):
                item = item.strip()
                if item:
                    result.append(item)

        return result

    def item_list(
        value: Any,
    ) -> list[AtsItem]:

        if not isinstance(value, list):
            return []

        result = []

        for item in value:

            if isinstance(item, dict):

                title = string(
                    item.get("title")
                    or item.get("name")
                    or item.get("issue")
                    or item.get("field")
                )

                explanation = string(
                    item.get("explanation")
                    or item.get("reason")
                    or item.get("description")
                    or item.get("details")
                )

                if title:
                    result.append(
                        AtsItem(
                            title=title,
                            explanation=explanation,
                        )
                    )

            elif isinstance(item, str):

                item = item.strip()

                if item:
                    result.append(
                        AtsItem(
                            title=item,
                            explanation="",
                        )
                    )

        return result

    def improvement_list(
        value: Any,
    ) -> list[AtsImprovement]:

        if not isinstance(value, list):
            return []

        result = []

        for item in value:

            if isinstance(item, dict):

                title = string(
                    item.get("title")
                    or item.get("name")
                    or item.get("action")
                )

                explanation = string(
                    item.get("explanation")
                    or item.get("reason")
                    or item.get("description")
                    or item.get("details")
                )

                priority = string(
                    item.get("priority"),
                    "Medium",
                )

                if title:
                    result.append(
                        AtsImprovement(
                            title=title,
                            explanation=explanation,
                            priority=priority,
                        )
                    )

            elif isinstance(item, str):

                item = item.strip()

                if item:
                    result.append(
                        AtsImprovement(
                            title=item,
                            explanation="",
                            priority="Medium",
                        )
                    )

        return result

    score = integer(
        data.get("ats_score"),
        0,
    )

    role_match = integer(
        data.get("role_match_score"),
        0,
    )

    role = string(
        data.get("role"),
        job_role.strip() or "Target Role",
    )

    score_label = string(
        data.get("score_label"),
    )

    if not score_label:
        if score >= 85:
            score_label = "Excellent"
        elif score >= 70:
            score_label = "Good"
        elif score >= 50:
            score_label = "Needs Improvement"
        else:
            score_label = "Weak"

    return AtsResult(
        ats_score=score,
        score_label=score_label,
        role_match_score=role_match,
        role=role,
        summary=string(
            data.get("summary"),
            "ATS analysis completed.",
        ),
        strengths=item_list(
            data.get("strengths"),
        ),
        missing_fields=item_list(
            data.get("missing_fields"),
        ),
        improvements=improvement_list(
            data.get("improvements"),
        ),
        keyword_gaps=string_list(
            data.get("keyword_gaps"),
        ),
        matched_keywords=string_list(
            data.get("matched_keywords"),
        ),
        formatting_issues=string_list(
            data.get("formatting_issues"),
        ),
        action_plan=string_list(
            data.get("action_plan"),
        ),
    )


def call_sarvam(
    resume_text: str,
    job_role: str,
    job_description: str,
    response_format: dict[str, Any],
) -> tuple[str, dict[str, Any]]:

    if not settings.sarvam_api_key:
        raise HTTPException(
            status_code=500,
            detail="SARVAM_API_KEY is not configured.",
        )

    payload = {
        "model": settings.sarvam_resume_model or "sarvam-105b",
        "messages": [
            {
                "role": "system",
                "content": (
                    "You are JOBIX's professional ATS engine. "
                    "Return only valid structured JSON."
                ),
            },
            {
                "role": "user",
                "content": build_prompt(
                    resume_text,
                    job_role,
                    job_description,
                ),
            },
        ],
        "temperature": 0.0,
        "reasoning_effort": None,
        "max_tokens": 4096,
        "response_format": response_format,
    }

    try:
        response = requests.post(
            "https://api.sarvam.ai/v1/chat/completions",
            headers={
                "api-subscription-key": settings.sarvam_api_key,
                "Content-Type": "application/json",
            },
            json=payload,
            timeout=75,
        )
    except requests.RequestException as exc:
        raise HTTPException(
            status_code=502,
            detail=f"Could not reach Sarvam: {exc}",
        )

    try:
        body = response.json()
    except ValueError:
        body = {}

    if response.status_code >= 400:

        error = (
            body.get("error", {})
            if isinstance(body, dict)
            else {}
        )

        message = (
            error.get("message")
            if isinstance(error, dict)
            else None
        )

        raise HTTPException(
            status_code=502,
            detail=(
                "Sarvam request failed: "
                f"{message or response.text[:500]}"
            ),
        )

    choices = (
        body.get("choices")
        if isinstance(body, dict)
        else None
    )

    if not isinstance(choices, list) or not choices:
        raise HTTPException(
            status_code=502,
            detail="Sarvam returned no choices for the ATS request.",
        )

    choice = (
        choices[0]
        if isinstance(choices[0], dict)
        else {}
    )

    message = (
        choice.get("message")
        if isinstance(choice.get("message"), dict)
        else {}
    )

    content = parse_content(
        message.get("content")
    )

    metadata = {
        "finish_reason": choice.get("finish_reason"),
        "model": body.get("model"),
        "usage": body.get("usage"),
    }

    return content, metadata


def analyze_with_sarvam(
    resume_text: str,
    job_role: str,
    job_description: str,
) -> AtsResult:
    schema_format = {
        "type": "json_schema",
        "json_schema": {
            "name": "jobix_ats_result",
            "strict": True,
            "schema": build_schema(),
        },
    }

    content, metadata = call_sarvam(
        resume_text,
        job_role,
        job_description,
        schema_format,
    )

    if content:
        try:
            data = parse_json_text(content)

            return normalize_result(
                data,
                job_role,
            )
        except Exception:
            pass

    fallback_format = {
        "type": "json_object",
    }

    fallback_content, fallback_metadata = call_sarvam(
        resume_text,
        job_role,
        job_description,
        fallback_format,
    )

    if fallback_content:
        try:
            data = parse_json_text(fallback_content)

            return normalize_result(
                data,
                job_role,
            )
        except Exception:
            pass

    finish_reason = (
        fallback_metadata.get("finish_reason")
        or metadata.get("finish_reason")
        or "unknown"
    )

    raise HTTPException(
        status_code=502,
        detail=(
            "JOBIX could not parse the ATS response from Sarvam. "
            f"finish_reason={finish_reason}."
        ),
    )


@router.post(
    "/check",
    response_model=AtsResult,
)
async def check_resume(
    resume: UploadFile = File(...),
    job_role: str = Form(""),
    job_description: str = Form(""),
    current_user: User = Depends(get_current_user),
):
    original_filename = resume.filename or "resume"
    original_content_type = resume.content_type or ""
    original_bytes = await resume.read()

    try:
        extracted_resume_text = await asyncio.to_thread(
            extract_uploaded_resume,
            original_filename,
            original_bytes,
            original_content_type,
        )
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=str(exc),
        )

    resume.file = BytesIO(
        extracted_resume_text.encode("utf-8")
    )
    resume.filename = "resume.txt"

    if not resume.filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Please upload a resume.",
        )

    file_data = await resume.read()

    if not file_data:
        raise HTTPException(
            status_code=400,
            detail="The uploaded resume is empty.",
        )

    if len(file_data) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=400,
            detail="Resume file must be 10 MB or smaller.",
        )

    resume.file = BytesIO(file_data)

    resume_text = extracted_resume_text

    jd = job_description.strip()[:MAX_JOB_DESCRIPTION]

    return analyze_with_sarvam(
        resume_text,
        job_role.strip(),
        jd,
    )
